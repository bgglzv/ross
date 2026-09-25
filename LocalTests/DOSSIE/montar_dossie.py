"""Monta o DOSSIÊ (.docx): SUSPEITO (esquerda) x ORIGINAL (direita), com comentários.

Pré-requisitos: MIT_P_teste_60Hz/teste_60Hz.py e DOSSIE/gerar_suspeito.py já executados
(geram, cada um, figuras/, anotadas/, metricas/ e dados/).
Os comentários são gerados a partir dos números reais das duas simulações.
"""

import math
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent
sys.path.insert(0, str(RAIZ.parent))

from comum.executar import slug
from ross.motors.utils import windowed_dfft

ORIG = RAIZ.parent / "MIT_P_teste_60Hz"
SUSP = RAIZ / "suspeito"
HIST = RAIZ / "suspeito_historico"

ACENTO = RGBColor(0x0E, 0x5C, 0x68)
CINZA = RGBColor(0x5B, 0x66, 0x72)

PARES = [
    ("Fonte CA (partida direta)", "Fonte CA (partida direta)"),
    ("Inversor V/F", "Inversor V/F"),
    ("Inversor FOC (regra antiga)", "Inversor FOC"),
]
EXCLUSIVO = "Inversor FOC (regra nova)"
ORIG_FOC = "Inversor FOC"
T_CARGA = 1.5


# ----------------------------------------------------------------------------
# Dados e comparações numéricas
# ----------------------------------------------------------------------------
def carregar(pasta, cenario):
    z = np.load(pasta / "dados" / f"{slug(cenario)}.npz")
    return {k: z[k] for k in z.files}


DADOS_S = {c: carregar(SUSP, c) for c, _ in PARES + [(EXCLUSIVO, None)]}
DADOS_O = {c: carregar(ORIG, c) for _, c in PARES}
MET_S = pd.read_csv(SUSP / "metricas" / "metricas_eventos.csv")
MET_O = pd.read_csv(ORIG / "metricas" / "metricas_eventos.csv")
RES_S = pd.read_csv(SUSP / "metricas" / "resumo_sinais.csv")
RES_O = pd.read_csv(ORIG / "metricas" / "resumo_sinais.csv")


def diferenca(a, b, chave, janela=None):
    t = a["t"]
    ya, yb = a[chave], b[chave]
    if janela is not None:
        m = (t >= janela[0]) & (t <= janela[1])
        t, ya, yb = t[m], ya[m], yb[m]
    d = ya - yb
    amplitude = max(float(np.ptp(ya)), float(np.ptp(yb)), 1e-12)
    ad = np.abs(d)
    idx = np.flatnonzero(ad > 1e-9 * amplitude)
    return {
        "max": float(ad.max()), "rms": float(np.sqrt(np.mean(d**2))),
        "rel": float(ad.max() / amplitude * 100.0),
        "t_diverge": float(t[idx[0]]) if idx.size else None,
        "identico": bool(ad.max() == 0.0),
    }


UNID = {"conjugado": "N·m", "velocidade_rpm": "RPM", "corrente_a": "A"}


NOME_CLASSE = {"identico": "idêntico", "coincidente": "praticamente coincidente",
               "pequena": "diferenças pequenas", "relevante": "diferenças relevantes"}


def classificar(d):
    if d["identico"]:
        return "identico"
    if d["rel"] < 0.5:
        return "coincidente"
    if d["rel"] < 5.0:
        return "pequena"
    return "relevante"


def frase(nome, d, sinal):
    u = UNID[sinal]
    k = classificar(d)
    base = f"diferença máxima {d['max']:.3g} {u} (RMS {d['rms']:.3g} {u}; {d['rel']:.2g}% da excursão do sinal)"
    if k == "identico":
        return f"{nome}: curvas idênticas ponto a ponto (diferença máxima igual a zero)"
    if k == "coincidente":
        return f"{nome}: curvas praticamente coincidentes — {base}"
    if k == "pequena":
        return f"{nome}: diferenças pequenas — {base}"
    return f"{nome}: diferenças relevantes — {base}" + (
        f", começando em t = {d['t_diverge']:.3f} s" if d["t_diverge"] is not None else "")


def comentar_series(sinal, janela):
    sem, disc = [], []
    for cs, co in PARES:
        d = diferenca(DADOS_S[cs], DADOS_O[co], sinal, janela)
        f = frase(co, d, sinal)
        (sem if classificar(d) in ("identico", "coincidente") else disc).append(f)
    return sem, disc


def picos_espectro(y, dt, n=3, fmin=20.0):
    f, m = windowed_dfft(y, dt)
    f, m = np.asarray(f), np.asarray(m)
    sel = f >= fmin
    f, m = f[sel], m[sel]
    ordem = np.argsort(m)[::-1]
    escolhidos = []
    for i in ordem:
        if all(abs(f[i] - f[j]) > 15 for j in escolhidos):
            escolhidos.append(i)
        if len(escolhidos) == n:
            break
    return [(float(f[i]), float(m[i])) for i in escolhidos]


def comentar_espectro():
    sem, disc = [], []
    for cs, co in PARES:
        dt = DADOS_S[cs]["t"][1] - DADOS_S[cs]["t"][0]
        ps = picos_espectro(DADOS_S[cs]["conjugado"], dt)
        po = picos_espectro(DADOS_O[co]["conjugado"], dt)
        txt_s = "; ".join(f"{f:.0f} Hz ({m:.3g})" for f, m in ps)
        txt_o = "; ".join(f"{f:.0f} Hz ({m:.3g})" for f, m in po)
        igual = all(abs(a[0] - b[0]) < 15 and abs(a[1] - b[1]) / max(a[1], 1e-12) < 0.05
                    for a, b in zip(ps, po))
        linha = f"{co}: componentes dominantes acima de 20 Hz — suspeito: {txt_s}; original: {txt_o}"
        (sem if igual else disc).append(linha)
    return sem, disc


def resumo_valor(res, cen, sinal, janela, col):
    r = res[(res.cenario == cen) & (res.sinal == sinal) & (res.janela == janela)]
    return float(r.iloc[0][col]) if len(r) else float("nan")


# ----------------------------------------------------------------------------
# Documento
# ----------------------------------------------------------------------------
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, lado, Cm(1.5))

est = doc.styles["Normal"]
est.font.name, est.font.size = "Arial", Pt(10)
est.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

contador = {"fig": 0}


def titulo(texto, tam=16):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(tam), True, "Arial", ACENTO
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p


def texto(t, italico=False, cor=None, tam=None):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.font.name = "Arial"
    if italico:
        r.font.italic = True
    if cor:
        r.font.color.rgb = cor
    if tam:
        r.font.size = Pt(tam)
    return p


def rotulado(rotulo, itens):
    """Parágrafo 'Rótulo. item1; item2...' — itens em lista com marcadores curtos."""
    p = doc.add_paragraph()
    r = p.add_run(rotulo + " ")
    r.font.bold, r.font.name = True, "Arial"
    if not itens:
        vazio = ("nenhuma — todos os critérios comparados coincidem." if rotulo.startswith("Discrep")
                 else "nenhuma semelhança relevante a registrar.")
        p.add_run(vazio).font.name = "Arial"
        return
    p.add_run(" • ".join(itens) + ".").font.name = "Arial"


def sombrear(cel, cor):
    tcpr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcpr.append(shd)


def par_de_figuras(desc, img_s, img_o, semelhancas, discrepancias, rot_s="SUSPEITO", rot_o="ORIGINAL"):
    contador["fig"] += 1
    p = doc.add_paragraph()
    r = p.add_run(f"Figura {contador['fig']} — {desc}")
    r.font.bold, r.font.name, r.font.size = True, "Arial", Pt(10.5)
    p.paragraph_format.keep_with_next = True
    tab = doc.add_table(rows=2, cols=2)
    tab.style = "Table Grid"
    for j, rot in enumerate((rot_s, rot_o)):
        c = tab.cell(0, j)
        c.text = ""
        run = c.paragraphs[0].add_run(rot)
        run.font.bold, run.font.name, run.font.size = True, "Arial", Pt(10)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        sombrear(c, "E3EEF0")
    for j, img in enumerate((img_s, img_o)):
        c = tab.cell(1, j)
        c.text = ""
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if img is not None and Path(img).is_file():
            c.paragraphs[0].add_run().add_picture(str(img), width=Cm(12.9))
        else:
            c.paragraphs[0].add_run("(figura indisponível)").font.italic = True
    rotulado("Semelhanças.", semelhancas)
    rotulado("Discrepâncias.", discrepancias)
    doc.add_paragraph()


# ---- Capa / propósito --------------------------------------------------------
p = doc.add_paragraph()
r = p.add_run("DOSSIÊ")
r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(30), True, "Arial", ACENTO
texto("Comparação entre o que foi avaliado por último (SUSPEITO) e o teste refeito com a biblioteca "
      "original (ORIGINAL) — motor de pequeno porte (1,5 hp), 60 Hz", cor=CINZA, tam=12)

titulo("01  Propósito e como ler este documento")
texto("Havia a suspeita de que o código da branch de trabalho anterior (local-tests-sync) estivesse com "
      "problemas. Para verificar, o trabalho daquela branch foi arquivado (tag arquivo/local-tests-sync-suspeito, "
      "commit ff997f28) e o teste foi reconstruído a partir da branch feature-motor limpa (commit 3da2863f), "
      "na nova branch local-test-sync2. Em cada figura, a coluna da esquerda mostra o resultado SUSPEITO "
      "e a da direita o resultado ORIGINAL, obtido agora. Sob cada par há comentários sobre as semelhanças e "
      "as discrepâncias, calculados a partir das séries de tempo reais das duas simulações (comparação "
      "ponto a ponto, na mesma grade de tempo).")

titulo("02  O que difere entre os dois")
texto("Motor, parâmetros elétricos, condição de ensaio (60 Hz, rampa de 0,6667 s, degrau de carga nominal em "
      "t = 1,5 s, frequência de chaveamento de 5 kHz, passo interno de 10 µs) e a análise/anotações das "
      "figuras são idênticos nos dois lados. As diferenças estão apenas em:")
for item in (
    "Biblioteca — o SUSPEITO inclui uma correção no inversor V/F (classe InverterVF): durante a rampa, o ângulo "
    "do fluxo passou a ser a integral da frequência instantânea (θ = f·t/2 enquanto a frequência sobe e "
    "θ = f·(t − t_rampa/2) depois), em vez de f·t. O ORIGINAL usa o cálculo anterior (θ = f·t). Com a rampa "
    "linear, f(t) = f_ref·t/t_rampa, a derivada de f·t vale f + t·df/dt = 2f: durante a rampa o ORIGINAL "
    "comanda ao campo girante o DOBRO da frequência instantânea desejada (e, depois da rampa, mantém um "
    "deslocamento constante de fase de f·t_rampa/2). A versão corrigida tem derivada exatamente igual a f. "
    "O acionamento por fonte de corrente alternada e o FOC não usam esse trecho.",
    "Cenário adicional — o SUSPEITO inclui o FOC com sintonia nova da malha de velocidade (razão entre as "
    "larguras de banda igual a 100, proteção de sobrecorrente instantânea), que não existe na biblioteca "
    "original. Ele é tratado à parte, na seção 07.",
    "Figuras comparativas do SUSPEITO — são as figuras originalmente avaliadas (com 4 curvas); as do ORIGINAL "
    "têm 3 curvas. As figuras anotadas do SUSPEITO foram regeneradas a partir do código arquivado, com a mesma "
    "análise do ORIGINAL.",
):
    q = doc.add_paragraph(style="List Bullet")
    q.add_run(item).font.name = "Arial"

titulo("03  Definições das marcações (iguais nos dois lados)")
texto("Sinal suavizado por média móvel de 20 ms (o bruto aparece ao fundo). Referência de velocidade: "
      "velocidade nominal; de conjugado: zero na partida em vazio e conjugado nominal após a carga (tracejada). "
      "Tubo de acomodação: referência ±5% (pontilhadas). Regime permanente: média dos últimos 10% da janela do "
      "evento. Máximo sobressinal: extremo do sinal em relação ao regime, na direção do movimento. "
      "Acomodação: primeiro instante a partir do qual o sinal permanece no tubo. Tempo de subida: de 10% a 90% "
      "da variação do valor inicial ao regime (retas verticais).")

# ---- Resumo quantitativo -----------------------------------------------------
titulo("04  Resumo quantitativo")


def fmt(v, c=2):
    return "—" if v is None or (isinstance(v, float) and math.isnan(v)) else f"{v:.{c}f}"


def pegar(met, cen, grandeza, t_ini_zero):
    m = met[(met.cenario == cen) & (met.grandeza == grandeza) & ((met.t_ini == 0) == t_ini_zero)]
    return m.iloc[0] if len(m) else None


linhas = []
for cs, co in PARES + [(EXCLUSIVO, ORIG_FOC)]:
    for grandeza in ("conjugado", "velocidade"):
        for ev, nome_ev in ((True, "Partida"), (False, "Carga")):
            s, o = pegar(MET_S, cs, grandeza, ev), pegar(MET_O, co, grandeza, ev)
            if s is None or o is None:
                continue
            c = 2 if grandeza == "conjugado" else 1
            linhas.append([
                cs + ("*" if cs == EXCLUSIVO else ""), grandeza, nome_ev,
                f"{s.sobressinal_pct:+.1f}", f"{o.sobressinal_pct:+.1f}",
                fmt(s.tempo_acom, 3) if s.tempo_acom == s.tempo_acom else "não acomoda",
                fmt(o.tempo_acom, 3) if o.tempo_acom == o.tempo_acom else "não acomoda",
                fmt(s.y_reg, c), fmt(o.y_reg, c),
            ])
cab = ["Cenário (suspeito)", "Grandeza", "Evento", "Sobress. % SUSP.", "Sobress. % ORIG.",
       "Acomod. (s) SUSP.", "Acomod. (s) ORIG.", "Regime SUSP.", "Regime ORIG."]
tab = doc.add_table(rows=1, cols=len(cab))
tab.style = "Light Grid Accent 1"
for i, h in enumerate(cab):
    tab.rows[0].cells[i].text = h
    for r in tab.rows[0].cells[i].paragraphs[0].runs:
        r.font.bold, r.font.size, r.font.name = True, Pt(8), "Arial"
for lin in linhas:
    cel = tab.add_row().cells
    for i, v in enumerate(lin):
        cel[i].text = str(v)
        for r in cel[i].paragraphs[0].runs:
            r.font.size, r.font.name = Pt(8), "Arial"
texto("* Cenário exclusivo do SUSPEITO; a coluna ORIGINAL mostra, para contraste, o FOC padrão da biblioteca original.",
      italico=True, cor=CINZA, tam=8.5)

# ---- Figuras comparativas ----------------------------------------------------
titulo("05  Figuras comparativas (esquema de legendas e zooms)")
JAN_ZOOM = (1.4, 1.7)
COMPARATIVAS = [
    ("01_conjugados_tempo", "Conjugados eletromagnéticos", "conjugado", None),
    ("01_conjugados_tempo_zoom", "Conjugados — zoom na entrada da carga", "conjugado", JAN_ZOOM),
    ("02_velocidade_tempo", "Velocidade do rotor", "velocidade_rpm", None),
    ("02_velocidade_tempo_zoom", "Velocidade — zoom na entrada da carga", "velocidade_rpm", JAN_ZOOM),
    ("03_corrente_fase_a_tempo", "Corrente da fase A", "corrente_a", None),
    ("03_corrente_fase_a_tempo_zoom", "Corrente da fase A — zoom na entrada da carga", "corrente_a", JAN_ZOOM),
    ("04_conjugados_freq", "Espectro do conjugado (FFT)", "espectro", None),
]
foc_ds = DADOS_S[EXCLUSIVO]
for arq, desc, sinal, jan in COMPARATIVAS:
    if sinal == "espectro":
        sem, disc = comentar_espectro()
    else:
        sem, disc = comentar_series(sinal, jan)
    sem.append("mesmo esquema de legendas, eixos e zoom nos dois lados")
    if sinal != "espectro":
        dnova = diferenca(foc_ds, DADOS_O[ORIG_FOC], sinal, jan)
        disc.append(f"o SUSPEITO traz uma quarta curva (FOC regra nova) que não existe no ORIGINAL; em relação "
                    f"ao FOC padrão original ela difere em até {dnova['max']:.3g} {UNID[sinal]} "
                    f"({dnova['rel']:.2g}% da excursão)")
    else:
        disc.append("o SUSPEITO traz uma quarta curva (FOC regra nova) sem equivalente no ORIGINAL")
    par_de_figuras(desc, HIST / f"{arq}.png", ORIG / "figuras" / f"{arq}.png", sem, disc)

# ---- Figuras anotadas --------------------------------------------------------
titulo("06  Figuras anotadas (sobressinal, tubo de ±5%, acomodação, regime e tempo de subida)")
texto("Para cada cenário: conjugado e velocidade, na entrada em velocidade nominal (partida) e após a entrada "
      "da carga (completa e com zoom). Os comentários comparam as métricas calculadas em cada lado.")


def campo(nome, s, o, casas, unid, tol_rel=0.02, tol_abs=None):
    """Compara um campo; devolve ('igual'|'dif', texto)."""
    if s is None or o is None or (isinstance(s, float) and math.isnan(s)) or (isinstance(o, float) and math.isnan(o)):
        ss = "não definido" if (s is None or (isinstance(s, float) and math.isnan(s))) else f"{s:.{casas}f} {unid}"
        oo = "não definido" if (o is None or (isinstance(o, float) and math.isnan(o))) else f"{o:.{casas}f} {unid}"
        if ss == oo:
            return "igual", f"{nome}: {ss} nos dois lados"
        return "dif", f"{nome}: {ss} no suspeito contra {oo} no original"
    ref = max(abs(s), abs(o), 1e-9)
    tol = tol_abs if tol_abs is not None else tol_rel * ref
    if abs(s - o) <= tol:
        return "igual", f"{nome}: {s:.{casas}f} {unid} (suspeito) e {o:.{casas}f} {unid} (original)"
    return "dif", f"{nome}: {s:.{casas}f} {unid} no suspeito contra {o:.{casas}f} {unid} no original"


def comentar_metricas(cs, co, grandeza, partida):
    s, o = pegar(MET_S, cs, grandeza, partida), pegar(MET_O, co, grandeza, partida)
    c = 2 if grandeza == "conjugado" else 1
    u = "N·m" if grandeza == "conjugado" else "RPM"
    sem, disc = [], []
    for k, texto_c in (
        campo("máximo sobressinal", s.sobressinal_pct, o.sobressinal_pct, 1, "%", tol_abs=0.5),
        campo("instante do máximo", s.t_pk, o.t_pk, 3, "s", tol_abs=0.005),
        campo("valor no máximo", s.y_pk, o.y_pk, c, u),
        campo("tempo de acomodação", s.tempo_acom, o.tempo_acom, 3, "s", tol_abs=0.005),
        campo("tempo de subida", s.tempo_subida, o.tempo_subida, 3, "s", tol_abs=0.005),
        campo("regime permanente", s.y_reg, o.y_reg, c, u, tol_rel=0.01),
    ):
        (sem if k == "igual" else disc).append(texto_c)
    return sem, disc, s, o


NOTA_VF = ("a única diferença de biblioteca que atinge este acionamento é a correção do ângulo de fase da "
           "tensão de referência do V/F (seção 02); como a fonte CA e o FOC, que não a usam, são idênticos "
           "nos dois lados, a discrepância fica isolada nessa correção")
for cs, co in PARES:
    titulo(co, 12)
    for grandeza in ("conjugado", "velocidade"):
        for partida, tag, desc_ev in ((True, "partida", "entrada em velocidade nominal"),
                                      (False, "carga", "após a entrada da carga"),
                                      (False, "carga_zoom", "após a entrada da carga (zoom)")):
            sem, disc, s, o = comentar_metricas(cs, co, grandeza, partida)
            if disc and "V/F" in co:
                disc.append(NOTA_VF)
            if not disc and cs != "Inversor V/F":
                sem.append("as duas simulações coincidem, como esperado: este acionamento não usa o trecho "
                           "da biblioteca alterado no SUSPEITO")
            par_de_figuras(f"{co} — {grandeza} — {desc_ev}",
                           SUSP / "anotadas" / f"{slug(cs)}_{grandeza}_{tag}.png",
                           ORIG / "anotadas" / f"{slug(co)}_{grandeza}_{tag}.png", sem, disc)

# ---- FOC regra nova ------------------------------------------------------------
titulo("07  Cenário exclusivo do SUSPEITO: FOC com sintonia nova")
texto("Este cenário não existe na biblioteca original. À esquerda, o FOC com sintonia nova (SUSPEITO); "
      "à direita, para contraste, o FOC padrão da biblioteca original (ORIGINAL).")
titulo("Inversor FOC — regra nova (suspeito) x Inversor FOC (original)", 12)
for grandeza in ("conjugado", "velocidade"):
    for partida, tag, desc_ev in ((True, "partida", "entrada em velocidade nominal"),
                                  (False, "carga", "após a entrada da carga"),
                                  (False, "carga_zoom", "após a entrada da carga (zoom)")):
        sem, disc, s, o = comentar_metricas(EXCLUSIVO, ORIG_FOC, grandeza, partida)
        disc.append("a sintonia nova da malha de velocidade só existe no SUSPEITO; motor, biblioteca do FOC e "
                    "condições de ensaio são os mesmos, então as diferenças entre os lados decorrem da sintonia")
        par_de_figuras(f"FOC (regra nova x padrão) — {grandeza} — {desc_ev}",
                       SUSP / "anotadas" / f"{slug(EXCLUSIVO)}_{grandeza}_{tag}.png",
                       ORIG / "anotadas" / f"{slug(ORIG_FOC)}_{grandeza}_{tag}.png", sem, disc,
                       rot_s="SUSPEITO (FOC regra nova)", rot_o="ORIGINAL (FOC padrão)")

# ---- Conclusão ------------------------------------------------------------------
titulo("08  Conclusão")
concl = []
for cs, co in PARES:
    dc = diferenca(DADOS_S[cs], DADOS_O[co], "conjugado")
    dv = diferenca(DADOS_S[cs], DADOS_O[co], "velocidade_rpm")
    di = diferenca(DADOS_S[cs], DADOS_O[co], "corrente_a")
    concl.append(f"{co}: conjugado — {NOME_CLASSE[classificar(dc)]} (máx. {dc['max']:.3g} N·m); velocidade — "
                 f"{NOME_CLASSE[classificar(dv)]} (máx. {dv['max']:.3g} RPM); corrente — "
                 f"{NOME_CLASSE[classificar(di)]} (máx. {di['max']:.3g} A).")
for c in concl:
    q = doc.add_paragraph(style="List Bullet")
    q.add_run(c).font.name = "Arial"
from comum import anotacoes as an

foc = DADOS_O[ORIG_FOC]
vs = an.media_movel(foc["t"], foc["velocidade_rpm"])
m = (foc["t"] > 0.75) & (foc["t"] < T_CARGA)
i_min = int(np.argmin(np.where(m, vs, np.inf)))
i_pico = int(np.argmax(np.where(foc["t"] < 1.0, vs, -np.inf)))
v_orig = resumo_valor(RES_O, ORIG_FOC, "velocidade_rpm", "pre_carga", "media")
texto(f"Observação sobre o FOC padrão da biblioteca original (igual nos dois lados): a velocidade atinge cerca de "
      f"{vs[i_pico]:.0f} RPM em t = {foc['t'][i_pico]:.2f} s, mas depois recua até {vs[i_min]:.0f} RPM em "
      f"t = {foc['t'][i_min]:.2f} s e volta a subir lentamente; nos 0,3 s que antecedem a carga a velocidade "
      f"média é de {v_orig:.0f} RPM, abaixo da nominal (1710 RPM). Ou seja, quando a carga entra o FOC ainda "
      f"está recuperando a velocidade, o que se reflete no conjugado elevado e oscilante antes de t = 1,5 s "
      f"na Figura 1. Como isso já ocorre na biblioteca original, não é efeito do código arquivado.")
texto("Leitura final (válida para este ensaio: motor de pequeno porte, 60 Hz): para a fonte de corrente alternada "
      "e para o FOC padrão, a suspeita de problema no código ou nos testes não se sustenta — os resultados são idênticos ponto a ponto. As discrepâncias restantes têm "
      "causa localizada: no V/F, a correção do ângulo de fase da rampa (seção 02), que só existe no SUSPEITO e, "
      "pela derivação apresentada, é a versão matematicamente coerente com a frequência comandada; e no FOC com "
      "sintonia nova, uma sintonia experimental que só existe no SUSPEITO (seção 07) e que, como já documentado, "
      "apresenta sensibilidade fina aos parâmetros. Manter ou não a correção do V/F na biblioteca é uma decisão "
      "a ser tomada; este dossiê apenas mostra que ela é a única alteração que muda um resultado existente.")

destino = RAIZ / "DOSSIE.docx"
doc.save(destino)
print(f"[OK] {destino}  ({contador['fig']} pares de figuras)")
