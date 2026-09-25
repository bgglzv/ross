"""Relatório (.docx) e apresentação (.pptx) de um teste de 60 Hz.

Texto corrido em português, sem menção a arquivos ou scripts; Arial, parágrafos
justificados, fundo branco e acento em verde-azulado (mesmo padrão já usado).
"""

import os

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Cm as PCm, Pt as PPt

from comum.executar import slug

ACENTO = RGBColor(0x0E, 0x5C, 0x68)
CINZA = RGBColor(0x5B, 0x66, 0x72)
TEXTO = RGBColor(0x1A, 0x20, 0x27)


def _f(v, casas=2, sufixo=""):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    return f"{v:.{casas}f}{sufixo}"


def _acom(v):
    return "não acomoda" if v is None or (isinstance(v, float) and np.isnan(v)) else f"{v:.3f}"


def linhas_tabela(metricas, grandeza):
    casas = 2 if grandeza == "conjugado" else 1
    linhas = []
    df = metricas[metricas.grandeza == grandeza]
    for _, r in df.iterrows():
        linhas.append([
            r.cenario, "Partida" if r.t_ini == 0 else "Carga",
            f"{r.sobressinal_pct:+.1f}", _acom(r.tempo_acom),
            _f(r.tempo_subida, 3) if r.tempo_subida == r.tempo_subida else "—",
            _f(r.y_reg, casas),
        ])
    return linhas


CABECALHO = ["Cenário", "Evento", "Máx. sobressinal (%)", "Acomodação ±5% (s)", "Tempo de subida (s)", "Regime"]


def _regra(doc):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    b = ppr.makeelement(qn("w:pBdr"), {})
    b.append(b.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "6",
                                            qn("w:space"): "1", qn("w:color"): "C3CAD1"}))
    ppr.append(b)


def _secao(doc, num, titulo):
    p = doc.add_paragraph()
    r = p.add_run(f"{num}  ")
    r.font.size, r.font.name, r.font.color.rgb = Pt(11), "Consolas", CINZA
    r2 = p.add_run(titulo)
    r2.font.size, r2.font.bold, r2.font.name, r2.font.color.rgb = Pt(16), True, "Arial", ACENTO
    p.paragraph_format.space_after = Pt(6)


def _tabela(doc, cab, linhas):
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(cab):
        t.rows[0].cells[i].text = h
        for r in t.rows[0].cells[i].paragraphs[0].runs:
            r.font.bold, r.font.size, r.font.name = True, Pt(8.5), "Arial"
    for lin in linhas:
        cel = t.add_row().cells
        for i, v in enumerate(lin):
            cel[i].text = str(v)
            for r in cel[i].paragraphs[0].runs:
                r.font.size, r.font.name = Pt(8.5), "Arial"


def _imagem(doc, caminho, legenda, largura=15.5):
    if os.path.isfile(caminho):
        doc.add_picture(caminho, width=Cm(largura))
        p = doc.add_paragraph(legenda)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size, r.font.italic, r.font.color.rgb, r.font.name = Pt(8.5), True, CINZA, "Arial"


def _parametros(motor, sim):
    rpm = float(motor.speed_nom) * 60 / (2 * np.pi)
    return [
        ["Potência nominal", f"{float(motor.power_nom) / 1000:.3g} kW"],
        ["Tensão nominal de fase", f"{float(motor.voltage_nom):.1f} V"],
        ["Frequência nominal", "60 Hz"],
        ["Polos", str(motor.n_poles)],
        ["Velocidade nominal", f"{rpm:.0f} RPM"],
        ["Conjugado nominal", f"{float(motor.Tnom):.2f} N·m"],
        ["Inércia do rotor", f"{float(motor.Ip_motor):.4g} kg·m²"],
        ["Frequência de chaveamento", "5000 Hz"],
        ["Rampa de aceleração", f"{sim.rampa:g} s"],
        ["Degrau de carga nominal em", f"t = {sim.t_carga:g} s"],
        ["Tempo total simulado", f"{sim.tf:g} s"],
    ]


def _leitura(metricas):
    d = metricas[(metricas.grandeza == "conjugado") & (metricas.t_ini > 0)]
    partes = []
    validos = d.dropna(subset=["tempo_acom"])
    if len(validos):
        m = validos.loc[validos.tempo_acom.idxmin()]
        partes.append(f"Após a entrada da carga, o conjugado do cenário “{m.cenario}” foi o que "
                      f"acomodou mais rápido no tubo de ±5% ({m.tempo_acom:.3f} s).")
    naoac = d[d.tempo_acom.isna()].cenario.tolist()
    if naoac:
        partes.append("Não acomodaram no tubo de ±5% do conjugado: " + ", ".join(naoac) + ".")
    v = metricas[(metricas.grandeza == "velocidade") & (metricas.t_ini > 0)]
    naov = v[v.tempo_acom.isna()].cenario.tolist()
    partes.append("Na velocidade, após a carga, " + (
        "não acomodaram no tubo de ±5%: " + ", ".join(naov) + "." if naov else
        "todos os cenários acomodaram no tubo de ±5%."))
    return " ".join(partes)


def gerar_relatorio_docx(pasta, nome_motor, motor, sim, metricas, resumo, arquivo):
    doc = Document()
    est = doc.styles["Normal"]
    est.font.name, est.font.size, est.font.color.rgb = "Arial", Pt(10.5), TEXTO
    est.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p = doc.add_paragraph()
    r = p.add_run("Teste a 60 Hz")
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(26), True, "Arial", ACENTO
    _regra(doc)
    p = doc.add_paragraph()
    r = p.add_run(f"{nome_motor} — partida direta, inversor V/F e inversor FOC")
    r.font.size, r.font.name, r.font.color.rgb = Pt(12), "Arial", CINZA
    doc.add_paragraph()

    _secao(doc, "01", "Objetivo e condições de ensaio")
    doc.add_paragraph(
        "Este relatório apresenta a resposta do motor a três formas de acionamento, todas em direção à "
        "frequência nominal de 60 Hz: fonte de corrente alternada ideal (partida direta), inversor com "
        "controle escalar V/F e inversor com controle vetorial por orientação de campo (FOC), este com a "
        "sintonia calculada automaticamente pela biblioteca. Após a acomodação em vazio, aplica-se um "
        "degrau de carga igual ao conjugado nominal do motor.")
    _tabela(doc, ["Parâmetro", "Valor"], _parametros(motor, sim))

    _secao(doc, "02", "Definições das marcações")
    doc.add_paragraph(
        "Todas as marcações são calculadas sobre o sinal suavizado por média móvel de 20 ms (o sinal bruto "
        "aparece ao fundo), para separar a dinâmica do acionamento da ondulação de chaveamento. A referência "
        "de velocidade é a velocidade nominal; a de conjugado é zero na partida em vazio e o conjugado nominal "
        "após a entrada da carga (reta tracejada). O tubo de acomodação é a referência ±5% (retas pontilhadas). "
        "O regime permanente é a média dos últimos 10% da janela de cada evento. O máximo sobressinal é o "
        "extremo do sinal, em relação ao regime, na direção do movimento; a acomodação é o primeiro instante "
        "a partir do qual o sinal permanece no tubo; e o tempo de subida é medido entre 10% e 90% da variação "
        "do valor inicial até o regime, com retas verticais marcando o início e o fim.")

    _secao(doc, "03", "Figuras comparativas")
    fig = os.path.join(pasta, "figuras")
    for arq, leg in (("01_conjugados_tempo", "Conjugados eletromagnéticos"),
                     ("01_conjugados_tempo_zoom", "Conjugados — zoom na entrada da carga"),
                     ("02_velocidade_tempo", "Velocidade do rotor"),
                     ("02_velocidade_tempo_zoom", "Velocidade — zoom na entrada da carga"),
                     ("03_corrente_fase_a_tempo", "Corrente da fase A"),
                     ("03_corrente_fase_a_tempo_zoom", "Corrente da fase A — zoom na entrada da carga"),
                     ("04_conjugados_freq", "Espectro do conjugado")):
        _imagem(doc, os.path.join(fig, arq + ".png"), leg)

    _secao(doc, "04", "Métricas por evento")
    doc.add_paragraph("Conjugado (N·m) — regime em N·m:")
    _tabela(doc, CABECALHO, linhas_tabela(metricas, "conjugado"))
    doc.add_paragraph()
    doc.add_paragraph("Velocidade — regime em RPM:")
    _tabela(doc, CABECALHO, linhas_tabela(metricas, "velocidade"))
    doc.add_paragraph()
    doc.add_paragraph(_leitura(metricas))

    _secao(doc, "05", "Figuras anotadas")
    ano = os.path.join(pasta, "anotadas")
    for cen in metricas.cenario.unique():
        p = doc.add_paragraph()
        r = p.add_run(cen)
        r.font.bold, r.font.name, r.font.color.rgb = True, "Arial", ACENTO
        for grandeza in ("conjugado", "velocidade"):
            for tag, leg in (("partida", "entrada em velocidade nominal"), ("carga", "após a entrada da carga"),
                             ("carga_zoom", "após a entrada da carga (zoom)")):
                _imagem(doc, os.path.join(ano, f"{slug(cen)}_{grandeza}_{tag}.png"),
                        f"{cen} — {grandeza} — {leg}", largura=15)

    doc.save(os.path.join(pasta, arquivo))
    print(f"[OK] {os.path.join(pasta, arquivo)}")


def _barra(slide, prs):
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, PCm(0.35))
    b.fill.solid()
    b.fill.fore_color.rgb = PRGB(0x0E, 0x5C, 0x68)
    b.line.fill.background()


def _titulo(slide, prs, texto):
    tb = slide.shapes.add_textbox(PCm(1), PCm(0.7), prs.slide_width - PCm(2), PCm(1.5))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = texto
    r.font.size, r.font.bold, r.font.name = PPt(24), True, "Arial"
    r.font.color.rgb = PRGB(0x0E, 0x5C, 0x68)


def _tabela_pptx(slide, cab, linhas, esq, topo, larg, alt):
    tab = slide.shapes.add_table(len(linhas) + 1, len(cab), esq, topo, larg, alt).table
    for j, h in enumerate(cab):
        c = tab.cell(0, j)
        c.text = h
        c.text_frame.paragraphs[0].runs[0].font.size = PPt(9)
        c.text_frame.paragraphs[0].runs[0].font.bold = True
        c.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        c.fill.solid()
        c.fill.fore_color.rgb = PRGB(0xE3, 0xEE, 0xF0)
    for i, lin in enumerate(linhas, start=1):
        for j, v in enumerate(lin):
            c = tab.cell(i, j)
            c.text = str(v)
            c.text_frame.paragraphs[0].runs[0].font.size = PPt(8.5)
            c.text_frame.paragraphs[0].runs[0].font.name = "Arial"
            c.fill.solid()
            c.fill.fore_color.rgb = PRGB(0xFF, 0xFF, 0xFF)


def gerar_apresentacao_pptx(pasta, nome_motor, motor, sim, metricas, arquivo):
    prs = Presentation()
    prs.slide_width, prs.slide_height = PCm(28.0), PCm(15.75)
    fig = os.path.join(pasta, "figuras")
    ano = os.path.join(pasta, "anotadas")
    vazio = prs.slide_layouts[6]

    s = prs.slides.add_slide(vazio)
    _barra(s, prs)
    _titulo(s, prs, f"Teste a 60 Hz — {nome_motor}")
    tb = s.shapes.add_textbox(PCm(1), PCm(2.6), PCm(26), PCm(3))
    tb.text_frame.word_wrap = True
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = ("Partida direta, inversor V/F e inversor FOC em direção à frequência nominal, "
              "seguidos de degrau de carga nominal.")
    r.font.size, r.font.name = PPt(14), "Arial"
    r.font.color.rgb = PRGB(0x5B, 0x66, 0x72)
    _tabela_pptx(s, ["Parâmetro", "Valor"], _parametros(motor, sim), PCm(1), PCm(5.2), PCm(14), PCm(9))

    def slide_imagens(titulo, arqs, pasta_img):
        sl = prs.slides.add_slide(vazio)
        _barra(sl, prs)
        _titulo(sl, prs, titulo)
        for k, a in enumerate(arqs):
            caminho = os.path.join(pasta_img, a + ".png")
            if os.path.isfile(caminho):
                sl.shapes.add_picture(caminho, PCm(0.7 + 13.4 * k), PCm(2.6), width=PCm(13.2))

    slide_imagens("Conjugado e velocidade", ["01_conjugados_tempo", "02_velocidade_tempo"], fig)
    slide_imagens("Zoom na entrada da carga", ["01_conjugados_tempo_zoom", "02_velocidade_tempo_zoom"], fig)

    sl = prs.slides.add_slide(vazio)
    _barra(sl, prs)
    _titulo(sl, prs, "Métricas após a entrada da carga")
    lin = [l for l in linhas_tabela(metricas, "conjugado") if l[1] == "Carga"]
    lin += [l for l in linhas_tabela(metricas, "velocidade") if l[1] == "Carga"]
    lin = [[("Conj. — " if i < len(lin) // 2 else "Vel. — ") + l[0]] + l[1:] for i, l in enumerate(lin)]
    _tabela_pptx(sl, CABECALHO, lin, PCm(0.7), PCm(2.6), PCm(26.6), PCm(8))

    for cen in metricas.cenario.unique():
        slide_imagens(f"{cen} — após a entrada da carga",
                      [f"{slug(cen)}_conjugado_carga_zoom", f"{slug(cen)}_velocidade_carga_zoom"], ano)

    prs.save(os.path.join(pasta, arquivo))
    print(f"[OK] {os.path.join(pasta, arquivo)}")
