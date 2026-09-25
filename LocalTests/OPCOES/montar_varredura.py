"""Documento da varredura do piso de fluxo (Opção 6): curvas sobrepostas e todos os índices."""

import json

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from plotly import graph_objects as go
from plotly.subplots import make_subplots

import barramento
import comum_opcoes as co
import opcao6
import opcao6_varredura as var
from comum import anotacoes as an
from comum import motores
from montar_opcao1 import ALTURA_CM, LARGURA_CM, _configurar_pagina

NOME_DOCUMENTO = "Opcao6-Enfraquecimento-de-campo_Varredura-do-fator.docx"
COR_ORIGINAL = "#555555"
CORES = {1.0: COR_ORIGINAL, 0.9: "#1f77b4", 0.8: "#2ca02c", 0.7: "#ff7f0e", 0.6: "#9467bd", 0.5: "#d62728"}


def _rotulo(k):
    return "Original (sem enfraquecimento)" if k == 1.0 else f"k_fw mínimo = {k:.1f}".replace(".", ",")


def _curto(k):
    return "orig." if k == 1.0 else f"k{k:.1f}".replace(".", ",")


def _arquivo(k):
    if k == 1.0:
        return barramento.DADOS / f"fator_{var.FATOR_BARRAMENTO}.npz", barramento.DADOS / f"fator_{var.FATOR_BARRAMENTO}.json"
    if k == 0.5:
        return opcao6.DADOS / f"fator_{var.FATOR_BARRAMENTO}.npz", opcao6.DADOS / f"fator_{var.FATOR_BARRAMENTO}.json"
    return var.DADOS / f"kmin_{k}.npz", var.DADOS / f"kmin_{k}.json"


def _carregar():
    curvas = {}
    for k in sorted(CORES, reverse=True):
        npz, js = _arquivo(k)
        d = np.load(npz)
        curvas[k] = dict(t=d["t"], vel=d["velocidade_rpm"], conj=d["conjugado"],
                         stats=json.loads(js.read_text()))
    return curvas


def _analises(curvas, motor):
    sim = motores.SIM_PEQUENO
    eventos = co._eventos(motor, sim.t_carga, sim.tf)
    for k, c in curvas.items():
        c["analise"] = {}
        c["suave"] = {}
        for grandeza, y in (("velocidade", c["vel"]), ("conjugado", c["conj"])):
            ys, an_ev = co._analise(c["t"], y, grandeza, eventos)
            c["suave"][grandeza] = ys
            c["analise"][grandeza] = an_ev
    return eventos


def figura_sobreposta(curvas, motor, titulo, janela, eventos_mostrados, com_rotulos):
    sim = motores.SIM_PEQUENO
    x0, x1 = janela
    fonte = 10 if com_rotulos else 9
    fig = make_subplots(rows=2, cols=1, shared_xaxes=True, vertical_spacing=0.06)
    linhas = ((1, "velocidade", "Velocidade (RPM)", (3, 1)), (2, "conjugado", "Conjugado (N·m)", (3, 2)))
    y_lim = {1: [], 2: []}
    ref_pintada = set()

    for row, grandeza, rotulo_y, casas in linhas:
        eixo, eixo_y = ("x", "y") if row == 1 else ("x2", "y2")
        for i, (k, c) in enumerate(curvas.items()):
            t, ys = c["t"], c["suave"][grandeza]
            m = (t >= x0) & (t <= x1)
            xs, yv = t[m], ys[m]
            y_lim[row] += [float(yv.min()), float(yv.max())]
            fig.add_trace(go.Scatter(x=xs, y=yv, name=_rotulo(k), legendgroup=str(k),
                                     showlegend=row == 1, line=dict(color=CORES[k], width=2.2)),
                          row=row, col=1)
            for ev_idx in eventos_mostrados:
                ev, res = c["analise"][grandeza][ev_idx]
                if (row, ev_idx) not in ref_pintada:
                    a, b = max(ev.t_ini, x0), min(ev.t_fim, x1)
                    fig.add_shape(type="line", x0=a, x1=b, y0=res["ref"], y1=res["ref"], row=row, col=1,
                                  line=dict(color="black", dash="dash", width=1.5))
                    for nivel in (res["tubo_sup"], res["tubo_inf"]):
                        fig.add_shape(type="line", x0=a, x1=b, y0=nivel, y1=nivel, row=row, col=1,
                                      line=dict(color="gray", dash="dot", width=1.5))
                    y_lim[row] += [res["tubo_inf"], res["tubo_sup"]]
                    ref_pintada.add((row, ev_idx))

                def ponto(x, y, simbolo, texto, ay):
                    if not (x0 <= x <= x1):
                        return
                    fig.add_trace(go.Scatter(x=[x], y=[y], mode="markers", showlegend=False, legendgroup=str(k),
                                             marker=dict(color=CORES[k], size=10, symbol=simbolo,
                                                         line=dict(color="black", width=1))),
                                  row=row, col=1)
                    if com_rotulos:
                        ax = 45 if (x - x0) / (x1 - x0) < 0.55 else -95
                        fig.add_annotation(x=x, y=y, xref=eixo, yref=eixo_y, text=texto, showarrow=True,
                                           arrowhead=2, arrowcolor=CORES[k], ax=ax, ay=ay,
                                           font=dict(size=fonte, color=CORES[k]),
                                           bgcolor="rgba(255,255,255,0.85)", bordercolor=CORES[k])

                unid = "N·m" if grandeza == "conjugado" else "RPM"
                lug = casas[1]
                passo = 26
                ponto(res["t_pk"], res["y_pk"], "circle",
                      f"{_curto(k)} pico {res['y_pk']:.{lug}f} {unid} ({res['sobressinal_pct']:+.0f}%) "
                      f"em {res['t_pk']:.3f} s", -30 - passo * i)
                if res["t_acom"] is not None:
                    ponto(res["t_acom"], res["y_acom"], "square",
                          f"{_curto(k)} acomod. {res['tempo_acom']:.3f} s ({res['y_acom']:.{lug}f} {unid})",
                          30 + passo * i)
                ponto(res["t_reg"], res["y_reg"], "diamond",
                      f"{_curto(k)} regime {res['y_reg']:.{lug}f} {unid}", -20 - passo * i)
                if res["t_sub0"] is not None:
                    for xv in (res["t_sub0"], res["t_sub1"]):
                        if x0 <= xv <= x1:
                            fig.add_vline(x=xv, row=row, col=1,
                                          line=dict(color=CORES[k], width=1.3, dash="dashdot"))
        pad = 0.15 * ((max(y_lim[row]) - min(y_lim[row])) or 1.0)
        fig.update_yaxes(range=[min(y_lim[row]) - pad, max(y_lim[row]) + pad], title_text=rotulo_y,
                         row=row, col=1)

    legendas = (("Referência", dict(color="black", dash="dash", width=1.5), "lines", None),
                ("Tubo ±5%", dict(color="gray", dash="dot", width=1.5), "lines", None),
                ("Subida 10–90% (linhas verticais)", dict(color="black", width=1.3, dash="dashdot"), "lines", None),
                ("Máx. sobressinal (círculo)", None, "markers", "circle"),
                ("Acomodação (quadrado)", None, "markers", "square"),
                ("Regime permanente (losango)", None, "markers", "diamond"))
    for nome, linha, modo, simbolo in legendas:
        if modo == "lines":
            fig.add_trace(go.Scatter(x=[None], y=[None], mode="lines", name=nome, line=linha), row=1, col=1)
        else:
            fig.add_trace(go.Scatter(x=[None], y=[None], mode="markers", name=nome,
                                     marker=dict(color="white", size=10, symbol=simbolo,
                                                 line=dict(color="black", width=1))), row=1, col=1)
    fig.update_xaxes(range=[x0, x1])
    fig.update_xaxes(title_text="Tempo (s)", row=2, col=1)
    fig.update_layout(title=dict(text=titulo, font=dict(size=18)), width=co.LARGURA_PX, height=co.ALTURA_PX,
                      legend=dict(orientation="h", yanchor="top", y=-0.07, x=0, font=dict(size=12)),
                      margin=dict(t=70, b=110, l=80, r=30), font=dict(family="Arial"))
    return fig


def _formatar(v, casas):
    return "—" if v is None else f"{v:.{casas}f}".replace(".", ",")


def _tabela_indices(doc, curvas, grandeza, ev_idx, titulo):
    lug = 2 if grandeza == "conjugado" else 1
    unid = "N·m" if grandeza == "conjugado" else "RPM"
    p = doc.add_paragraph()
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(10)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    colunas = ("Curva", "Sobressinal (%)", f"Pico ({unid})", "Instante do pico (s)", "Acomodação (s)",
               f"Valor na acomodação ({unid})", f"Regime ({unid})", "Subida 10–90% (s)", "Início → fim da subida (s)")
    tabela = doc.add_table(rows=1, cols=len(colunas))
    tabela.style = "Light Grid Accent 1"
    for c, nome in zip(tabela.rows[0].cells, colunas):
        c.text = nome
    for k, c in curvas.items():
        _, res = c["analise"][grandeza][ev_idx]
        subida = "—" if res["t_sub0"] is None else f"{res['t_sub0']:.3f} → {res['t_sub1']:.3f}".replace(".", ",")
        vals = (_rotulo(k), _formatar(res["sobressinal_pct"], 1), _formatar(res["y_pk"], lug),
                _formatar(res["t_pk"], 3), _formatar(res["tempo_acom"], 3) if res["t_acom"] is not None else "não acomoda",
                _formatar(res["y_acom"], lug), _formatar(res["y_reg"], lug),
                _formatar(res["tempo_subida"], 3), subida)
        for cel, v in zip(tabela.add_row().cells, vals):
            cel.text = v
    for row in tabela.rows:
        for cel in row.cells:
            for par in cel.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(8)
                    run.font.name = "Arial"


def _tabela_saturacao(doc, curvas):
    p = doc.add_paragraph()
    r = p.add_run("Saturação da referência de corrente q e fluxo aplicado (barramento 1,35 × V_linha)")
    r.bold = True
    r.font.size = Pt(10)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    colunas = ("Curva", "Saturado na rampa (%)", "Saturado em vazio (%)", "Saturado após a carga (%)",
               "Conj. mín. (N·m)", "Conj. máx. (N·m)", "Corrente de pico (A)", "k_fw médio em vazio", "k_fw final")
    tabela = doc.add_table(rows=1, cols=len(colunas))
    tabela.style = "Light Grid Accent 1"
    for c, nome in zip(tabela.rows[0].cells, colunas):
        c.text = nome
    for k, c in curvas.items():
        s = c["stats"]
        vals = (_rotulo(k), _formatar(s["pct_rampa"], 1), _formatar(s["pct_regime_vazio"], 1),
                _formatar(s["pct_apos_carga"], 1), _formatar(s["conjugado_minimo"], 2),
                _formatar(s["conjugado_maximo"], 2), _formatar(s["corrente_pico_a"], 2),
                _formatar(s.get("kfw_medio_vazio", 1.0), 2), _formatar(s.get("kfw_final", 1.0), 2))
        for cel, v in zip(tabela.add_row().cells, vals):
            cel.text = v
    for row in tabela.rows:
        for cel in row.cells:
            for par in cel.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(8)
                    run.font.name = "Arial"


def montar():
    sim = motores.SIM_PEQUENO
    motor = motores.motor_pequeno()
    var.FIGURAS.mkdir(exist_ok=True)
    curvas = _carregar()
    _analises(curvas, motor)

    janelas_carga = [co.janelas_de_zoom(c["t"], c["vel"], c["conj"], motor, sim.t_carga, sim.tf, sim.rampa)[1]
                     for c in curvas.values()]
    zoom_carga = (min(j[0] for j in janelas_carga), max(j[1] for j in janelas_carga))
    pedidos = (
        ("completa", "visão completa", (0.0, sim.tf), [0, 1], False),
        ("zoom_velocidade_nominal", "zoom na entrada em velocidade nominal",
         (max(0.0, sim.rampa - 0.25), sim.t_carga - 0.02), [0], True),
        ("zoom_carga", "zoom na entrada da carga", zoom_carga, [1], True),
    )
    doc = Document()
    _configurar_pagina(doc)
    for k, (chave, descricao, janela, eventos, rotulos) in enumerate(pedidos):
        if k:
            doc.add_page_break()
        titulo = f"Opção 6 — Varredura do piso do fator de fluxo (barramento 1,35 × V_linha) — {descricao}"
        fig = figura_sobreposta(curvas, motor, titulo, janela, eventos, rotulos)
        caminho = var.FIGURAS / f"varredura_{chave}.png"
        fig.write_image(str(caminho), width=co.LARGURA_PX, height=co.ALTURA_PX, scale=1.5)
        print(f"  [OK] {caminho}")
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(str(caminho), width=Cm(LARGURA_CM), height=Cm(ALTURA_CM))

    doc.add_page_break()
    titulo = doc.add_paragraph()
    r = titulo.add_run("Índices de desempenho por curva")
    r.bold = True
    r.font.size = Pt(13)
    for grandeza, ev_idx, nome in (("velocidade", 0, "Velocidade — entrada em velocidade nominal"),
                                   ("conjugado", 0, "Conjugado — entrada em velocidade nominal"),
                                   ("velocidade", 1, "Velocidade — após a entrada da carga"),
                                   ("conjugado", 1, "Conjugado — após a entrada da carga")):
        if nome.startswith("Velocidade — após"):
            doc.add_page_break()
        _tabela_indices(doc, curvas, grandeza, ev_idx, nome)
    _tabela_saturacao(doc, curvas)
    destino = var.PASTA / NOME_DOCUMENTO
    doc.save(destino)
    print(f"[OK] {destino}")
