"""Monta o documento único da Opção 3, com os limites de 1,5× e 2,0× lado a lado."""

import json

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

import barramento
import comum_opcoes as co
import opcao3
from comum import motores
from montar_opcao2 import (ALTURA_FIG_CM, ALTURA_FIG_PX, LARGURA_FIG_CM, LARGURA_FIG_PX,
                           _faixas, _tabela_resumo)
from montar_opcao1 import _configurar_pagina

NOME_DOCUMENTO = "Opcao3-Limite-corrente-inversor_15x-20x.docx"


def _rotulo(multiplo):
    return f"{multiplo:.2g}".replace(".", ",")


def _carregar(multiplo):
    d = np.load(opcao3.DADOS / f"multiplo_{multiplo}.npz")
    return d["t"], d["velocidade_rpm"], d["conjugado"]


def _figuras_do_par(motor, multiplos, chave, descricao, janela, eventos):
    sim = motores.SIM_PEQUENO
    dados = {m: _carregar(m) for m in multiplos}
    caminhos = []
    for m in multiplos:
        outro = next(o for o in multiplos if o != m)
        t, vel, conj = dados[m]
        titulo = f"Opção 3 — Limite de corrente {_rotulo(m)} × Is_nom — {descricao}"
        fig = co.figura_foc(t, vel, conj, motor, sim.t_carga, sim.tf, titulo, janela, eventos,
                            largura_px=LARGURA_FIG_PX, altura_px=ALTURA_FIG_PX,
                            y_extra=_faixas(*dados[outro], janela), fonte=9)
        fig.update_layout(title=dict(font=dict(size=13)), legend=dict(font=dict(size=10), y=-0.05),
                          margin=dict(t=60, b=150, l=70, r=20))
        caminho = opcao3.FIGURAS / f"multiplo_{m}_{chave}.png"
        fig.write_image(str(caminho), width=LARGURA_FIG_PX, height=ALTURA_FIG_PX, scale=2)
        print(f"  [OK] {caminho}")
        caminhos.append(caminho)
    return caminhos


def _janela_de_zoom_da_carga(motor, multiplos):
    sim = motores.SIM_PEQUENO
    janelas = []
    for m in multiplos:
        t, vel, conj = _carregar(m)
        janelas.append(co.janelas_de_zoom(t, vel, conj, motor, sim.t_carga, sim.tf, sim.rampa)[1])
    return min(j[0] for j in janelas), max(j[1] for j in janelas)


def montar():
    sim = motores.SIM_PEQUENO
    motor = motores.motor_pequeno()
    opcao3.FIGURAS.mkdir(exist_ok=True)
    multiplos = opcao3.MULTIPLOS
    pedidos = (
        ("completa", "visão completa", (0.0, sim.tf), [0, 1]),
        ("zoom_velocidade_nominal", "zoom na entrada em velocidade nominal",
         (max(0.0, sim.rampa - 0.25), sim.t_carga - 0.02), [0]),
        ("zoom_carga", "zoom na entrada da carga", _janela_de_zoom_da_carga(motor, multiplos), [1]),
    )
    doc = Document()
    _configurar_pagina(doc)
    for k, (chave, descricao, janela, eventos) in enumerate(pedidos):
        caminhos = _figuras_do_par(motor, multiplos, chave, descricao, janela, eventos)
        if k:
            doc.add_page_break()
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for caminho in caminhos:
            par.add_run().add_picture(str(caminho), width=Cm(LARGURA_FIG_CM), height=Cm(ALTURA_FIG_CM))

    linhas = [("Original (limite 3,0 × Is_nom)",
               json.loads((barramento.DADOS / "fator_1.35.json").read_text()))]
    for m in multiplos:
        linhas.append((f"Opção 3 (limite {_rotulo(m)} × Is_nom)",
                       json.loads((opcao3.DADOS / f"multiplo_{m}.json").read_text())))
    _tabela_resumo(doc, linhas, titulo="Opção 3 — Resumo da saturação (FOC, 60 Hz, motor pequeno, "
                                       "rampa de 0,6667 s, barramento padrão 1,35 × V_linha)")
    destino = opcao3.PASTA / NOME_DOCUMENTO
    doc.save(destino)
    print(f"[OK] {destino}")
