"""Documento A4 paisagem com duas variantes lado a lado (velocidade em cima, conjugado embaixo)."""

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import comum_opcoes as co
from comum import motores
from montar_opcao1 import _configurar_pagina
from montar_opcao2 import (ALTURA_FIG_CM, ALTURA_FIG_PX, LARGURA_FIG_CM, LARGURA_FIG_PX,
                           _faixas, _tabela_resumo)


def _pagina_de_texto(doc, titulo, secoes):
    doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    for cabecalho, itens in secoes:
        p = doc.add_paragraph()
        r = p.add_run(cabecalho)
        r.bold = True
        r.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        for item in itens:
            par = doc.add_paragraph(style="List Bullet")
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = par.add_run(item)
            run.font.size = Pt(10)
            run.font.name = "Arial"
            par.paragraph_format.space_after = Pt(3)


def montar_documento(variantes, pasta_figuras, titulo_figura, arquivo_dados, linhas_tabela,
                     titulo_tabela, destino, texto=None):
    """`variantes` é uma lista de duas chaves; `arquivo_dados(chave)` devolve o .npz da variante;
    `titulo_figura(chave, descricao)` devolve o título da figura."""
    sim = motores.SIM_PEQUENO
    motor = motores.motor_pequeno()
    pasta_figuras.mkdir(parents=True, exist_ok=True)
    dados = {}
    for v in variantes:
        d = np.load(arquivo_dados(v))
        dados[v] = (d["t"], d["velocidade_rpm"], d["conjugado"])

    janelas_carga = [co.janelas_de_zoom(*dados[v], motor, sim.t_carga, sim.tf, sim.rampa)[1]
                     for v in variantes]
    zoom_carga = (min(j[0] for j in janelas_carga), max(j[1] for j in janelas_carga))
    pedidos = (
        ("completa", "visão completa", (0.0, sim.tf), [0, 1]),
        ("zoom_velocidade_nominal", "zoom na entrada em velocidade nominal",
         (max(0.0, sim.rampa - 0.25), sim.t_carga - 0.02), [0]),
        ("zoom_carga", "zoom na entrada da carga", zoom_carga, [1]),
    )

    doc = Document()
    _configurar_pagina(doc)
    for k, (chave, descricao, janela, eventos) in enumerate(pedidos):
        if k:
            doc.add_page_break()
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for v in variantes:
            outro = next(o for o in variantes if o != v)
            t, vel, conj = dados[v]
            fig = co.figura_foc(t, vel, conj, motor, sim.t_carga, sim.tf, titulo_figura(v, descricao),
                                janela, eventos, largura_px=LARGURA_FIG_PX, altura_px=ALTURA_FIG_PX,
                                y_extra=_faixas(*dados[outro], janela), fonte=9)
            fig.update_layout(title=dict(font=dict(size=13)), legend=dict(font=dict(size=10), y=-0.05),
                              margin=dict(t=60, b=150, l=70, r=20))
            caminho = pasta_figuras / f"{v}_{chave}.png"
            fig.write_image(str(caminho), width=LARGURA_FIG_PX, height=ALTURA_FIG_PX, scale=2)
            print(f"  [OK] {caminho}")
            par.add_run().add_picture(str(caminho), width=Cm(LARGURA_FIG_CM), height=Cm(ALTURA_FIG_CM))
    if texto:
        _pagina_de_texto(doc, *texto)
    _tabela_resumo(doc, linhas_tabela, titulo=titulo_tabela)
    doc.save(destino)
    print(f"[OK] {destino}")
