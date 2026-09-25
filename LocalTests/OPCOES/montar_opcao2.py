"""Monta o documento único da Opção 2, com as figuras dos dois barramentos lado a lado."""

import json

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import barramento
import comum_opcoes as co
import opcao2
from comum import anotacoes as an
from comum import motores
from montar_opcao1 import _configurar_pagina

NOME_DOCUMENTO = "Opcao2-Limitar-modulo-vetor_CC-135x1414.docx"
LARGURA_FIG_PX = 900
ALTURA_FIG_PX = 1240
LARGURA_FIG_CM = 13.7
ALTURA_FIG_CM = LARGURA_FIG_CM * ALTURA_FIG_PX / LARGURA_FIG_PX


def _rotulo_fator(fator):
    return f"{fator:.4g}".replace(".", ",")


def _carregar(fator):
    d = np.load(opcao2.DADOS / f"fator_{fator}.npz")
    return d["t"], d["velocidade_rpm"], d["conjugado"]


def _faixas(t, vel, conj, janela):
    x0, x1 = janela
    m = (t >= x0) & (t <= x1)
    vs, cs = an.media_movel(t, vel)[m], an.media_movel(t, conj)[m]
    return {1: [float(vs.min()), float(vs.max())], 2: [float(cs.min()), float(cs.max())]}


def _figura_do_par(motor, fatores, chave, descricao, janela, eventos, arquivo):
    sim = motores.SIM_PEQUENO
    dados = {f: _carregar(f) for f in fatores}
    caminhos = []
    for f in fatores:
        outro = next(o for o in fatores if o != f)
        t, vel, conj = dados[f]
        titulo = f"Opção 2 — Barramento {_rotulo_fator(f)} × V_linha — {descricao}"
        fig = co.figura_foc(t, vel, conj, motor, sim.t_carga, sim.tf, titulo, janela, eventos,
                            largura_px=LARGURA_FIG_PX, altura_px=ALTURA_FIG_PX,
                            y_extra=_faixas(*dados[outro], janela), fonte=9)
        fig.update_layout(title=dict(font=dict(size=13)), legend=dict(font=dict(size=10), y=-0.05),
                          margin=dict(t=60, b=150, l=70, r=20))
        caminho = opcao2.FIGURAS / f"fator_{f}_{chave}.png"
        fig.write_image(str(caminho), width=LARGURA_FIG_PX, height=ALTURA_FIG_PX, scale=2)
        print(f"  [OK] {caminho}")
        caminhos.append(caminho)
    return caminhos


def _janela_de_zoom_da_carga(motor, fatores):
    sim = motores.SIM_PEQUENO
    janelas = []
    for f in fatores:
        t, vel, conj = _carregar(f)
        janelas.append(co.janelas_de_zoom(t, vel, conj, motor, sim.t_carga, sim.tf, sim.rampa)[1])
    return min(j[0] for j in janelas), max(j[1] for j in janelas)


def _tabela_resumo(doc, linhas, titulo="Opção 2 — Resumo da saturação (FOC, 60 Hz, motor pequeno, rampa de 0,6667 s)"):
    doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(13)
    p.paragraph_format.space_after = Pt(8)

    colunas = ("Variante", "Barramento (× V_linha)", "Saturado na rampa (%)", "Saturado em vazio (%)",
               "Saturado após a carga (%)", "Razão máx. |iq_ref| / limite", "Conjugado mín. (N·m)",
               "Conjugado máx. (N·m)", "Corrente de pico (A)")
    tabela = doc.add_table(rows=1, cols=len(colunas))
    tabela.style = "Light Grid Accent 1"
    for c, nome in zip(tabela.rows[0].cells, colunas):
        c.text = nome
    for variante, e in linhas:
        vals = (variante, f"{e['fator']:.4g}", f"{e['pct_rampa']:.1f}", f"{e['pct_regime_vazio']:.1f}",
                f"{e['pct_apos_carga']:.1f}", f"{e['razao_maxima']:.1f}", f"{e['conjugado_minimo']:.2f}",
                f"{e['conjugado_maximo']:.2f}", f"{e['corrente_pico_a']:.2f}")
        for c, v in zip(tabela.add_row().cells, vals):
            c.text = v.replace(".", ",")
    for row in tabela.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Arial"


def montar():
    sim = motores.SIM_PEQUENO
    motor = motores.motor_pequeno()
    opcao2.FIGURAS.mkdir(exist_ok=True)
    fatores = opcao2.FATORES
    pedidos = (
        ("completa", "visão completa", (0.0, sim.tf), [0, 1]),
        ("zoom_velocidade_nominal", "zoom na entrada em velocidade nominal",
         (max(0.0, sim.rampa - 0.25), sim.t_carga - 0.02), [0]),
        ("zoom_carga", "zoom na entrada da carga", _janela_de_zoom_da_carga(motor, fatores), [1]),
    )
    doc = Document()
    _configurar_pagina(doc)
    for k, (chave, descricao, janela, eventos) in enumerate(pedidos):
        caminhos = _figura_do_par(motor, fatores, chave, descricao, janela, eventos, None)
        if k:
            doc.add_page_break()
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for caminho in caminhos:
            par.add_run().add_picture(str(caminho), width=Cm(LARGURA_FIG_CM), height=Cm(ALTURA_FIG_CM))

    linhas = []
    for f in fatores:
        linhas.append(("Original (só limite em iq)", json.loads((barramento.DADOS / f"fator_{f}.json").read_text())))
        linhas.append(("Opção 2 (módulo do vetor)", json.loads((opcao2.DADOS / f"fator_{f}.json").read_text())))
    _tabela_resumo(doc, linhas)
    destino = opcao2.PASTA / NOME_DOCUMENTO
    doc.save(destino)
    print(f"[OK] {destino}")
