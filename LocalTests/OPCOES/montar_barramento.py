"""Monta as figuras e o documento do teste do barramento CC."""

import json

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import barramento
import comum_opcoes as co
from comum import motores
from montar_opcao1 import ALTURA_CM, LARGURA_CM, _configurar_pagina

NOME_DOCUMENTO = "Teste-Barramento-CC-folgado.docx"


def _rotulo_fator(fator):
    return f"{fator:.4g}".replace(".", ",")


def _gerar_figuras(motor, fator):
    sim = motores.SIM_PEQUENO
    dados = np.load(barramento.DADOS / f"fator_{fator}.npz")
    t, vel, conj = dados["t"], dados["velocidade_rpm"], dados["conjugado"]
    zoom_partida, zoom_carga = co.janelas_de_zoom(t, vel, conj, motor, sim.t_carga, sim.tf, sim.rampa)
    zoom_partida = (max(0.0, sim.rampa - 0.25), sim.t_carga - 0.02)
    pedidos = (
        ("completa", "visão completa", (0.0, sim.tf), [0, 1]),
        ("zoom_velocidade_nominal", "zoom na entrada em velocidade nominal", zoom_partida, [0]),
        ("zoom_carga", "zoom na entrada da carga", zoom_carga, [1]),
    )
    caminhos = []
    for chave, descricao, janela, eventos in pedidos:
        titulo = (f"Barramento CC = {_rotulo_fator(fator)} × V_linha — FOC, 60 Hz, "
                  f"rampa de 0,6667 s — {descricao}")
        fig = co.figura_foc(t, vel, conj, motor, sim.t_carga, sim.tf, titulo, janela, eventos)
        caminho = barramento.FIGURAS / f"fator_{fator}_{chave}.png"
        fig.write_image(str(caminho), width=co.LARGURA_PX, height=co.ALTURA_PX, scale=1.5)
        print(f"  [OK] {caminho}")
        caminhos.append(caminho)
    return caminhos


def _tabela_resumo(doc, estatisticas):
    doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run("Teste do barramento CC — Resumo (FOC, 60 Hz, motor pequeno, rampa de 0,6667 s)")
    r.bold = True
    r.font.size = Pt(13)
    p.paragraph_format.space_after = Pt(8)

    colunas = ("Barramento (× V_linha)", "Saturado na rampa (%)", "Saturado em vazio (%)",
               "Saturado após a carga (%)", "Razão máx. |iq_ref| / limite", "Conjugado mín. (N·m)",
               "Conjugado máx. (N·m)", "Corrente de pico (A)")
    tabela = doc.add_table(rows=1, cols=len(colunas))
    tabela.style = "Light Grid Accent 1"
    for c, nome in zip(tabela.rows[0].cells, colunas):
        c.text = nome
    for e in estatisticas:
        vals = (f"{e['fator']:.4g}", f"{e['pct_rampa']:.1f}", f"{e['pct_regime_vazio']:.1f}",
                f"{e['pct_apos_carga']:.1f}", f"{e['razao_maxima']:.1f}", f"{e['conjugado_minimo']:.2f}",
                f"{e['conjugado_maximo']:.2f}", f"{e['corrente_pico_a']:.2f}")
        for c, v in zip(tabela.add_row().cells, vals):
            c.text = v.replace(".", ",")
    for row in tabela.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Arial"


def montar():
    motor = motores.motor_pequeno()
    barramento.FIGURAS.mkdir(exist_ok=True)
    doc = Document()
    _configurar_pagina(doc)
    estatisticas = []
    primeira = True
    for fator in barramento.FATORES:
        arq = barramento.DADOS / f"fator_{fator}.json"
        if not arq.exists():
            print(f"Fator {fator} ainda não simulado; ignorado.")
            continue
        estatisticas.append(json.loads(arq.read_text()))
        for caminho in _gerar_figuras(motor, fator):
            if not primeira:
                doc.add_page_break()
            primeira = False
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(str(caminho), width=Cm(LARGURA_CM), height=Cm(ALTURA_CM))
    _tabela_resumo(doc, estatisticas)
    destino = barramento.PASTA / NOME_DOCUMENTO
    doc.save(destino)
    print(f"[OK] {destino}")
