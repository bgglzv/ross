"""Monta as figuras e o documento da Opção 1 a partir das rampas já simuladas."""

import json

import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import comum_opcoes as co
import opcao1
from comum import motores

NOME_DOCUMENTO = "Opcao1-Reduzir-a-demanda-de-corrente-na-origem.docx"
LARGURA_CM = 27.6
ALTURA_CM = LARGURA_CM * co.ALTURA_PX / co.LARGURA_PX


def _rotulo_rampa(rampa):
    return f"{rampa:.4g} s".replace(".", ",")


def _gerar_figuras(motor, rampa):
    dados = np.load(opcao1.DADOS / f"rampa_{rampa}.npz")
    t, vel, conj = dados["t"], dados["velocidade_rpm"], dados["conjugado"]
    t_carga, tf = opcao1.parametros(rampa)
    zoom_partida, zoom_carga = co.janelas_de_zoom(t, vel, conj, motor, t_carga, tf, rampa)
    zoom_partida = (max(0.0, rampa - 0.25), t_carga - 0.02)
    pedidos = (
        ("completa", "visão completa", (0.0, tf), [0, 1]),
        ("zoom_velocidade_nominal", "zoom na entrada em velocidade nominal", zoom_partida, [0]),
        ("zoom_carga", "zoom na entrada da carga", zoom_carga, [1]),
    )
    caminhos = []
    for chave, descricao, janela, eventos in pedidos:
        titulo = f"Opção 1 — FOC, rampa de {_rotulo_rampa(rampa)} — {descricao}"
        fig = co.figura_foc(t, vel, conj, motor, t_carga, tf, titulo, janela, eventos)
        caminho = opcao1.FIGURAS / f"rampa_{rampa}_{chave}.png"
        fig.write_image(str(caminho), width=co.LARGURA_PX, height=co.ALTURA_PX, scale=1.5)
        print(f"  [OK] {caminho}")
        caminhos.append(caminho)
    return caminhos


def _configurar_pagina(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for lado in ("left_margin", "right_margin"):
        setattr(sec, lado, Cm(1.05))
    sec.top_margin = sec.bottom_margin = Cm(0.8)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(10)
    estilo.paragraph_format.space_after = Pt(0)
    estilo.paragraph_format.space_before = Pt(0)


def _tabela_resumo(doc, estatisticas):
    doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run("Opção 1 — Resumo da saturação da referência de corrente q (FOC, 60 Hz, motor pequeno)")
    r.bold = True
    r.font.size = Pt(13)
    p.paragraph_format.space_after = Pt(8)

    colunas = ("Rampa (s)", "Saturado na rampa (%)", "Saturado em vazio (%)", "Saturado após a carga (%)",
               "Razão máx. |iq_ref| / limite", "Conjugado mín. (N·m)", "Conjugado máx. (N·m)", "Corrente de pico (A)")
    tabela = doc.add_table(rows=1, cols=len(colunas))
    tabela.style = "Light Grid Accent 1"
    for c, nome in zip(tabela.rows[0].cells, colunas):
        c.text = nome
    for e in estatisticas:
        vals = (f"{e['rampa']:.4g}", f"{e['pct_rampa']:.1f}", f"{e['pct_regime_vazio']:.1f}",
                f"{e['pct_apos_carga']:.1f}", f"{e['razao_maxima']:.1f}", f"{e['conjugado_minimo']:.2f}",
                f"{e['conjugado_maximo']:.2f}", f"{e['corrente_pico_a']:.2f}")
        for c, v in zip(tabela.add_row().cells, vals):
            c.text = v.replace(".", ",")
    for row in tabela.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Arial"


def montar():
    motor = motores.motor_pequeno()
    opcao1.FIGURAS.mkdir(exist_ok=True)
    doc = Document()
    _configurar_pagina(doc)
    estatisticas = []
    primeira = True
    for rampa in opcao1.RAMPAS:
        arq = opcao1.DADOS / f"rampa_{rampa}.json"
        if not arq.exists():
            print(f"Rampa {rampa} ainda não simulada; ignorada.")
            continue
        estatisticas.append(json.loads(arq.read_text()))
        for caminho in _gerar_figuras(motor, rampa):
            if not primeira:
                doc.add_page_break()
            primeira = False
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(str(caminho), width=Cm(LARGURA_CM), height=Cm(ALTURA_CM))
    _tabela_resumo(doc, estatisticas)
    destino = opcao1.PASTA / NOME_DOCUMENTO
    doc.save(destino)
    print(f"[OK] {destino}")
